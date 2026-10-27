from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import unicodedata

from docx import Document
from pypdf import PdfReader
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}


class DocumentProcessingError(ValueError):
    """Raised when a document cannot produce usable text."""


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str


@dataclass(frozen=True)
class TextChunk:
    page_number: int
    text: str


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def extract_text(content: bytes, filename: str) -> list[ExtractedPage]:
    if not content:
        raise DocumentProcessingError("Document is empty.")

    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentProcessingError(
            "Unsupported document type. Allowed types: PDF, DOCX, PPTX, TXT."
        )

    try:
        if extension == ".pdf":
            pages = _extract_pdf(content)
        elif extension == ".docx":
            pages = _extract_docx(content)
        elif extension == ".pptx":
            pages = _extract_pptx(content)
        else:
            pages = _extract_txt(content)
    except (DocumentProcessingError, UnicodeDecodeError):
        raise
    except Exception as exc:
        raise DocumentProcessingError(
            f"Could not read {extension.removeprefix('.').upper()} document."
        ) from exc

    cleaned = [
        ExtractedPage(page.page_number, clean_text(page.text))
        for page in pages
        if clean_text(page.text)
    ]
    if not cleaned:
        raise DocumentProcessingError("Document contains no extractable text.")
    return cleaned


def chunk_pages(
    pages: list[ExtractedPage],
    *,
    chunk_size: int,
    overlap: int,
) -> list[TextChunk]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive.")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size.")

    chunks: list[TextChunk] = []
    for page in pages:
        page_text = clean_text(page.text)
        start = 0
        while start < len(page_text):
            end = min(start + chunk_size, len(page_text))
            if end < len(page_text):
                boundary = page_text.rfind(" ", start, end + 1)
                if boundary > start:
                    end = boundary
            text = page_text[start:end].strip()
            if text:
                chunks.append(TextChunk(page.page_number, text))
            if end >= len(page_text):
                break
            next_start = max(end - overlap, start + 1)
            if overlap:
                boundary = page_text.find(" ", next_start, end)
                if boundary != -1:
                    next_start = boundary + 1
            start = next_start
    return chunks


def _extract_pdf(content: bytes) -> list[ExtractedPage]:
    reader = PdfReader(BytesIO(content))
    return [
        ExtractedPage(index, page.extract_text() or "")
        for index, page in enumerate(reader.pages, start=1)
    ]


def _extract_docx(content: bytes) -> list[ExtractedPage]:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return [ExtractedPage(1, "\n".join(parts))]


def _extract_pptx(content: bytes) -> list[ExtractedPage]:
    presentation = Presentation(BytesIO(content))
    pages = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                parts.append(shape.text)
        pages.append(ExtractedPage(number, "\n".join(parts)))
    return pages


def _extract_txt(content: bytes) -> list[ExtractedPage]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError("TXT document must be valid UTF-8.") from exc
    return [ExtractedPage(1, text)]
